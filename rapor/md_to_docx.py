"""
md_to_docx.py
--------------
STAJ_RAPORU.md dosyasini, staj yazim kurallarina uygun bicimlendirilmis
bir Word (.docx) belgesine cevirir.

SABLON UZERINE YAZAR
--------------------
Belge SIFIRDAN olusturulmaz; fakultenin resmi sablonu (sablon_rapor.docx)
acilir ve YALNIZCA govde icerigi degistirilir. Boylece her sayfada
tekrarlanan su ogeler oldugu gibi korunur:

  - Ust bilgi: okul logosu, "MUHENDISLIK MIMARLIK VE TASARIM FAKULTESI /
    STAJ RAPORU", dokuman no (RP-023), yayin/revizyon tarihleri, sayfa no
  - Alt bilgi: "Ogrenci - Ad Soyad / Imza", "Yetkili - Ad Soyad /
    Kase-Imza", "Hazirlayan BKK / Onaylayan KASGEM"
  - Sayfa boyutu ve kenar bosluklari (ust 3,5 cm — ust bilgi icin pay)

UYGULANAN KURALLAR (staj yonergesinden)
---------------------------------------
  - A4 sayfa, Times New Roman
  - Govde metni 11 punto, 1,5 satir araligi, IKI YANA yasli
  - Basliklar sola yasli, buyuk punto (12 pt kalin), asili girinti 0,76 cm
  - Paragraflar ve bolumler arasi 8 nk bosluk
  - Baslik ile kendi ilk paragrafi arasinda bosluk YOK
  - Tablo ve sekiller ORTALI
  - Tablo aciklamasi tablonun USTUNDE, sekil aciklamasi seklin ALTINDA
  - Sayfa numarasi sablonun ust bilgisindeki alandan gelir

KULLANIM
--------
    python rapor/md_to_docx.py

Cikti: rapor/STAJ_RAPORU.docx
"""

import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

KOK = os.path.dirname(os.path.abspath(__file__))
GIRDI = os.path.join(KOK, "STAJ_RAPORU.md")
CIKTI = os.path.join(KOK, "STAJ_RAPORU.docx")
SABLON = os.path.join(KOK, "sablon_rapor.docx")

YAZI_TIPI = "Times New Roman"
GOVDE_PUNTO = Pt(11)
BASLIK_PUNTO = Pt(12)
PARAGRAF_BOSLUGU = Pt(8)          # 8 nk
BASLIK_ASILI_GIRINTI = Cm(0.76)
SAYFA_GENISLIGI_CM = 17.0          # sablon kenar bosluklarina gore (21 - 2 - 2)


# --------------------------------------------------------------------------
# Belge kurulumu
# --------------------------------------------------------------------------


def belge_olustur():
    """Resmi sablonu acar ve YALNIZCA govdesini bosaltir.

    Ust/alt bilgi (logo, fakulte adi, dokuman no, imza alanlari) ve sayfa
    duzeni sablonda tanimlidir; bunlara DOKUNULMAZ. Govdedeki `sectPr`
    ogesi de korunur — bu oge, ust/alt bilgi baglantilarini ve sayfa
    boyutunu tasir; silinirse her sayfada tekrarlanan basliklar kaybolur.
    """
    if not os.path.exists(SABLON):
        sys.exit(
            f"Sablon bulunamadi: {SABLON}\n"
            "Fakultenin resmi RP-023 dosyasini bu ada kopyalayin."
        )
    belge = Document(SABLON)

    govde = belge.element.body
    for cocuk in list(govde):
        if cocuk.tag.endswith("}sectPr"):
            continue
        govde.remove(cocuk)

    normal = belge.styles["Normal"]
    normal.font.name = YAZI_TIPI
    normal.font.size = GOVDE_PUNTO
    if normal.element.rPr is not None and normal.element.rPr.rFonts is not None:
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), YAZI_TIPI)

    return belge


def paragraf_bicimle(p, hizalama=WD_ALIGN_PARAGRAPH.JUSTIFY, bosluk_once=Pt(0),
                     bosluk_sonra=PARAGRAF_BOSLUGU, satir_araligi=1.5):
    bf = p.paragraph_format
    bf.alignment = hizalama
    bf.space_before = bosluk_once
    bf.space_after = bosluk_sonra
    bf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    bf.line_spacing = satir_araligi
    return p


# --------------------------------------------------------------------------
# Markdown ayristirma
# --------------------------------------------------------------------------

KALIN = re.compile(r"\*\*(.+?)\*\*")
KOD = re.compile(r"`([^`]+)`")


def metin_ekle(p, metin):
    """**kalin** ve `kod` isaretlemelerini Word bicimlendirmesine cevirir."""
    parcalar = re.split(r"(\*\*.+?\*\*|`[^`]+`)", metin)
    for parca in parcalar:
        if not parca:
            continue
        if parca.startswith("**") and parca.endswith("**"):
            c = p.add_run(parca[2:-2])
            c.bold = True
        elif parca.startswith("`") and parca.endswith("`"):
            c = p.add_run(parca[1:-1])
            c.font.name = "Consolas"
            c.font.size = Pt(10)
        else:
            c = p.add_run(parca)
        c.font.name = c.font.name or YAZI_TIPI


def baslik_ekle(belge, metin):
    p = belge.add_paragraph()
    paragraf_bicimle(p, WD_ALIGN_PARAGRAPH.LEFT,
                     bosluk_once=PARAGRAF_BOSLUGU, bosluk_sonra=Pt(0))
    # Asili girinti 0,76 cm
    p.paragraph_format.left_indent = BASLIK_ASILI_GIRINTI
    p.paragraph_format.first_line_indent = -BASLIK_ASILI_GIRINTI
    p.paragraph_format.keep_with_next = True
    c = p.add_run(metin)
    c.bold = True
    c.font.name = YAZI_TIPI
    c.font.size = BASLIK_PUNTO
    return p


def alt_baslik_ekle(belge, metin):
    p = belge.add_paragraph()
    paragraf_bicimle(p, WD_ALIGN_PARAGRAPH.LEFT,
                     bosluk_once=PARAGRAF_BOSLUGU, bosluk_sonra=Pt(0))
    p.paragraph_format.keep_with_next = True
    c = p.add_run(metin)
    c.bold = True
    c.italic = True
    c.font.name = YAZI_TIPI
    c.font.size = GOVDE_PUNTO
    return p


def govde_ekle(belge, metin):
    p = belge.add_paragraph()
    paragraf_bicimle(p)
    metin_ekle(p, metin)
    return p


def aciklama_ekle(belge, metin, ustte):
    """Tablo/sekil aciklamasi: ortali, kalin etiket."""
    p = belge.add_paragraph()
    paragraf_bicimle(p, WD_ALIGN_PARAGRAPH.CENTER,
                     bosluk_once=Pt(0) if ustte else Pt(4),
                     bosluk_sonra=Pt(4) if ustte else PARAGRAF_BOSLUGU)
    p.paragraph_format.keep_with_next = ustte
    metin_ekle(p, metin)
    return p


def gorsel_ekle(belge, yol, genislik_cm=None):
    if not os.path.exists(yol):
        print(f"  UYARI: gorsel bulunamadi: {yol}")
        return
    p = belge.add_paragraph()
    paragraf_bicimle(p, WD_ALIGN_PARAGRAPH.CENTER,
                     bosluk_once=PARAGRAF_BOSLUGU, bosluk_sonra=Pt(2))
    p.paragraph_format.keep_with_next = True
    c = p.add_run()
    c.add_picture(yol, width=Cm(genislik_cm or SAYFA_GENISLIGI_CM))


def liste_ekle(belge, metin, sirali=False, sira_no=None):
    """Madde isaretli / numarali liste satiri.

    Sablonda "List Bullet" ve "List Number" stilleri bulunmadigi icin
    isaret elle eklenir; boylece belge sablonun stil listesinden bagimsiz
    kalir."""
    p = belge.add_paragraph()
    paragraf_bicimle(p, WD_ALIGN_PARAGRAPH.JUSTIFY, bosluk_sonra=Pt(4))
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    isaret = f"{sira_no}. " if sirali and sira_no else ("" if sirali else "• ")
    if isaret:
        c = p.add_run(isaret)
        c.font.name = YAZI_TIPI
        c.font.size = GOVDE_PUNTO
    metin_ekle(p, metin)
    return p


def kod_blogu_ekle(belge, satirlar):
    p = belge.add_paragraph()
    paragraf_bicimle(p, WD_ALIGN_PARAGRAPH.LEFT,
                     bosluk_once=Pt(4), bosluk_sonra=PARAGRAF_BOSLUGU,
                     satir_araligi=1.0)
    p.paragraph_format.left_indent = Cm(1.0)
    c = p.add_run("\n".join(satirlar))
    c.font.name = "Consolas"
    c.font.size = Pt(9.5)
    c.font.color.rgb = RGBColor(0x30, 0x30, 0x30)


def tablo_kenarliklari(tablo):
    """Tabloya ince siyah kenarlik uygular.

    Resmi sablonda "Table Grid" stili tanimli degildir (yalnizca "Normal
    Table" vardir), bu yuzden kenarliklar dogrudan XML ile veriliyor.
    Boylece sablonun stil listesine bagimli kalinmiyor."""
    ozellikler = tablo._tbl.tblPr
    kenarliklar = OxmlElement("w:tblBorders")
    for kenar in ("top", "left", "bottom", "right", "insideH", "insideV"):
        oge = OxmlElement(f"w:{kenar}")
        oge.set(qn("w:val"), "single")
        oge.set(qn("w:sz"), "4")          # 4 = 0,5 punto
        oge.set(qn("w:space"), "0")
        oge.set(qn("w:color"), "000000")
        kenarliklar.append(oge)
    ozellikler.append(kenarliklar)


def tablo_ekle(belge, satirlar):
    """Markdown tablosunu Word tablosuna cevirir."""
    hucreler = [[h.strip() for h in s.strip().strip("|").split("|")]
                for s in satirlar]
    baslik = hucreler[0]
    govde = hucreler[2:]          # 1. satir ayirici (---|---)

    tablo = belge.add_table(rows=1, cols=len(baslik))
    tablo.alignment = WD_TABLE_ALIGNMENT.CENTER
    tablo_kenarliklari(tablo)

    for i, metin in enumerate(baslik):
        hucre = tablo.rows[0].cells[i]
        hucre.text = ""
        p = hucre.paragraphs[0]
        paragraf_bicimle(p, WD_ALIGN_PARAGRAPH.CENTER,
                         bosluk_sonra=Pt(0), satir_araligi=1.0)
        c = p.add_run(metin.replace("**", ""))
        c.bold = True
        c.font.name = YAZI_TIPI
        c.font.size = Pt(10)

    for satir in govde:
        yeni = tablo.add_row().cells
        for i, metin in enumerate(satir[:len(baslik)]):
            yeni[i].text = ""
            p = yeni[i].paragraphs[0]
            paragraf_bicimle(p, WD_ALIGN_PARAGRAPH.LEFT,
                             bosluk_sonra=Pt(0), satir_araligi=1.0)
            metin_ekle(p, metin)
            for c in p.runs:
                c.font.size = Pt(10)

    bos = belge.add_paragraph()
    paragraf_bicimle(bos, bosluk_sonra=PARAGRAF_BOSLUGU, satir_araligi=1.0)
    bos.runs.clear() if bos.runs else None
    return tablo


# --------------------------------------------------------------------------
# Ana donusum
# --------------------------------------------------------------------------


def donustur():
    with open(GIRDI, encoding="utf-8") as f:
        satirlar = f.read().split("\n")

    belge = belge_olustur()
    i = 0
    sayac = {"sekil": 0, "tablo": 0}
    bekleyen_aciklama = None        # tablo aciklamasi tablodan ONCE gelir

    while i < len(satirlar):
        satir = satirlar[i]
        cikplak = satir.strip()

        # --- bos satir / yatay cizgi ---
        if not cikplak or cikplak == "---":
            i += 1
            continue

        # --- kod blogu ---
        if cikplak.startswith("```"):
            i += 1
            blok = []
            while i < len(satirlar) and not satirlar[i].strip().startswith("```"):
                blok.append(satirlar[i])
                i += 1
            kod_blogu_ekle(belge, blok)
            i += 1
            continue

        # --- basliklar ---
        if cikplak.startswith("# "):
            p = belge.add_paragraph()
            paragraf_bicimle(p, WD_ALIGN_PARAGRAPH.CENTER,
                             bosluk_sonra=Pt(12))
            c = p.add_run(cikplak[2:])
            c.bold = True
            c.font.name = YAZI_TIPI
            c.font.size = Pt(16)
            i += 1
            continue
        if cikplak.startswith("## "):
            baslik_ekle(belge, cikplak[3:])
            i += 1
            continue
        if cikplak.startswith("### "):
            alt_baslik_ekle(belge, cikplak[4:])
            i += 1
            continue

        # --- gorsel ---
        eslesme = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", cikplak)
        if eslesme:
            yol = os.path.join(KOK, eslesme.group(2))
            # Dar gorseller sayfa genisligine zorlanmasin
            genislik = SAYFA_GENISLIGI_CM
            try:
                from struct import unpack

                with open(yol, "rb") as g:
                    g.read(16)
                    w, h = unpack(">II", g.read(8))
                if w < 700:
                    genislik = min(SAYFA_GENISLIGI_CM, w / 96 * 2.54 * 1.6)
                # Cok uzun gorselleri sayfaya sigdir (en fazla ~19 cm yukseklik)
                oran = h / w
                if genislik * oran > 19.0:
                    genislik = 19.0 / oran
            except Exception:
                pass
            gorsel_ekle(belge, yol, genislik)
            i += 1
            continue

        # --- tablo / sekil aciklamasi ---
        if cikplak.startswith("**Tablo "):
            sayac["tablo"] += 1
            bekleyen_aciklama = cikplak
            i += 1
            continue
        if cikplak.startswith("**Şekil "):
            sayac["sekil"] += 1
            aciklama_ekle(belge, cikplak, ustte=False)
            i += 1
            continue

        # --- tablo ---
        if cikplak.startswith("|"):
            blok = []
            while i < len(satirlar) and satirlar[i].strip().startswith("|"):
                blok.append(satirlar[i])
                i += 1
            if bekleyen_aciklama:
                aciklama_ekle(belge, bekleyen_aciklama, ustte=True)
                bekleyen_aciklama = None
            tablo_ekle(belge, blok)
            continue

        # --- alinti (doldurulacak notlar) ---
        if cikplak.startswith("> "):
            blok = []
            while i < len(satirlar) and satirlar[i].strip().startswith(">"):
                blok.append(satirlar[i].strip().lstrip(">").strip())
                i += 1
            # Doldurulacak notlar: kirmizi italik, **kalin** isaretlemesi
            # duz metin olarak kalmasin diye metin_ekle ile islenir.
            for j, metin in enumerate([x for x in blok if x]):
                p = belge.add_paragraph()
                paragraf_bicimle(
                    p, WD_ALIGN_PARAGRAPH.LEFT,
                    bosluk_once=PARAGRAF_BOSLUGU if j == 0 else Pt(0),
                    bosluk_sonra=Pt(3),
                )
                p.paragraph_format.left_indent = Cm(1.0)
                metin_ekle(p, metin)
                for c in p.runs:
                    c.italic = True
                    c.font.size = Pt(10.5)
                    c.font.color.rgb = RGBColor(0x80, 0x00, 0x00)
            continue

        # --- listeler ---
        if cikplak.startswith("- "):
            liste_ekle(belge, cikplak[2:])
            i += 1
            continue
        eslesme_no = re.match(r"^(\d+)\. (.*)", cikplak)
        if eslesme_no:
            liste_ekle(belge, eslesme_no.group(2), sirali=True,
                       sira_no=eslesme_no.group(1))
            i += 1
            continue

        # --- normal paragraf ---
        govde_ekle(belge, cikplak)
        i += 1

    belge.save(CIKTI)
    bolum = belge.sections[0]
    print(f"  Olusturuldu: {CIKTI}")
    print(f"  {sayac['sekil']} sekil · {sayac['tablo']} tablo")
    print(f"  ust bilgi tablosu : {len(bolum.header.tables)} (logo ve fakulte basligi)")
    print(f"  alt bilgi tablosu : {len(bolum.footer.tables)} (imza alanlari)")
    print(f"  sayfa duzeni      : {bolum.page_width.cm:.0f}x{bolum.page_height.cm:.0f} cm, "
          f"ust kenar {bolum.top_margin.cm:.1f} cm")


if __name__ == "__main__":
    if not os.path.exists(GIRDI):
        sys.exit(f"Girdi bulunamadi: {GIRDI}")
    donustur()
